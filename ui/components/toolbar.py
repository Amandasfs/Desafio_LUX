import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook

def create_toolbar(root, dashboard):
    toolbar = tk.Frame(root, bg="#053c9b", height=50)
    toolbar.pack(fill="x")

    search_container = tk.Frame(toolbar, bg="#00215b")
    search_container.pack(side="left", padx=10)

    tk.Label(
        search_container,
        text="🔍",
        bg="#00215b",
        fg="white",
        font=("Arial", 12, "bold"),
    ).pack(side="left")

    search_entry = tk.Entry(search_container, width=30, font=("Arial", 11))
    search_entry.pack(side="left", padx=5, pady=10)

    def get_search_term():
        return search_entry.get().lower().strip()

    dashboard.get_search_term = get_search_term

    tk.Button(
        search_container,
        text="Buscar",
        bg="#004aad",
        fg="white",
        font=("Arial", 10, "bold"),
        relief="flat",
        command=dashboard.buscar,
        cursor="hand2"
    ).pack(side="left", padx=5)

    tk.Button(
        search_container,
        text="Organizar",
        bg="#0066cc",
        fg="white",
        font=("Arial", 10, "bold"),
        relief="flat",
        command=dashboard.organizar_dados,
        cursor="hand2"
    ).pack(side="left", padx=5)

    # Botão para cadastrar novo contato
    tk.Button(
        search_container,
        text="Cadastrar Contato",
        bg="#00aaff",
        fg="white",
        font=("Arial", 10, "bold"),
        relief="flat",
        command=lambda: abrir_formulario_contato(dashboard),
        cursor="hand2"
    ).pack(side="left", padx=5)

    # Botão para alterar contato
    tk.Button(
        search_container,
        text="Alterar Contato",
        bg="#ffaa00",
        fg="white",
        font=("Arial", 10, "bold"),
        relief="flat",
        command=abrir_alterar_contato,
        cursor="hand2"
    ).pack(side="left", padx=5)

    return toolbar

# Função que abre a janela de cadastro
def abrir_formulario_contato(dashboard):
    form = tk.Toplevel()
    form.title("📋 Cadastrar Novo Contato")
    form.geometry("750x700")
    form.configure(bg="#f8fafc")

    # Scrollbar
    canvas = tk.Canvas(form, bg="#f8fafc", highlightthickness=0)
    scrollbar = ttk.Scrollbar(form, orient="vertical", command=canvas.yview)
    scroll_frame = ttk.Frame(canvas)

    scroll_frame.bind(
        "<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
    )

    canvas.create_window((0, 0), window=scroll_frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    # Estilo geral
    estilo_titulo = ("Arial", 13, "bold")
    estilo_label = ("Arial", 10)
    bg_card = "#ffffff"

    # === SEÇÃO CLIENTE ===
    cliente_frame = tk.LabelFrame(scroll_frame, text="👤 Dados do Cliente", font=estilo_titulo,
                                  bg=bg_card, fg="#00215b", padx=10, pady=10, labelanchor="n")
    cliente_frame.pack(fill="x", padx=20, pady=(20, 10))

    campos_cliente = [
        "Identificador", "Nome", "Cargo", "Empresa",
        "Telefone", "E-mail", "Gestor Responsável (LUX)"
    ]

    entradas = {}

    for campo in campos_cliente:
        linha = tk.Frame(cliente_frame, bg=bg_card)
        linha.pack(fill="x", pady=5)
        tk.Label(linha, text=campo, bg=bg_card, font=estilo_label, width=25, anchor="w").pack(side="left")
        entrada = ttk.Entry(linha, width=40)
        entrada.pack(side="left", padx=10)
        entradas[campo] = entrada

    # === SEÇÃO EMPRESA ===
    empresa_frame = tk.LabelFrame(scroll_frame, text="🏢 Dados da Empresa", font=estilo_titulo,
                                  bg=bg_card, fg="#00215b", padx=10, pady=10, labelanchor="n")
    empresa_frame.pack(fill="x", padx=20, pady=(10, 20))

    campos_empresa = [
        "Endereço - Rua", "Endereço - Numero", "Endereço - Estado",
        "Endereço - Cidade", "Endereço - CEP", "Razão Social", "CNPJ",
        "Distribuidora", "Modalidade Tarifária", "Consumo Ponta (kWh)",
        "Consumo Fora Ponta (kWh)", "Valor Médio da Fatura (R$)"
    ]

    for campo in campos_empresa:
        linha = tk.Frame(empresa_frame, bg=bg_card)
        linha.pack(fill="x", pady=5)
        tk.Label(linha, text=campo, bg=bg_card, font=estilo_label, width=25, anchor="w").pack(side="left")
        entrada = ttk.Entry(linha, width=40)
        entrada.pack(side="left", padx=10)
        entradas[campo] = entrada

    # === BOTÕES ===
    botoes_frame = tk.Frame(scroll_frame, bg="#f8fafc")
    botoes_frame.pack(pady=30)

    def salvar_contato():
        valores = [entradas[c].get() for c in (campos_cliente + campos_empresa)]
        try:
            wb = load_workbook("dados_brutos.xlsx")
            ws = wb.active
            ws.append(valores)
            wb.save("dados_brutos.xlsx")
            messagebox.showinfo("✅ Sucesso", "Contato cadastrado com sucesso!")
            form.destroy()
        except Exception as e:
            messagebox.showerror("❌ Erro", f"Não foi possível salvar o contato.\n{e}")

    ttk.Button(botoes_frame, text="💾 Salvar Contato", command=salvar_contato).pack(side="left", padx=10)
    ttk.Button(botoes_frame, text="❌ Cancelar", command=form.destroy).pack(side="left", padx=10)

# Função que abre a janela de alteração
def abrir_alterar_contato():
    form = tk.Toplevel()
    form.title("Alterar Contato")
    form.geometry("400x250")

    tk.Label(form, text="Nome do Contato:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
    nome_entry = tk.Entry(form, width=30)
    nome_entry.grid(row=0, column=1, padx=10, pady=10)

    tk.Label(form, text="Telefone:").grid(row=1, column=0, padx=10, pady=10, sticky="w")
    telefone_entry = tk.Entry(form, width=30)
    telefone_entry.grid(row=1, column=1, padx=10, pady=10)

    tk.Label(form, text="E-mail:").grid(row=2, column=0, padx=10, pady=10, sticky="w")
    email_entry = tk.Entry(form, width=30)
    email_entry.grid(row=2, column=1, padx=10, pady=10)

    def buscar_contato():
        nome = nome_entry.get().strip().lower()
        try:
            wb = load_workbook("dados_brutos.xlsx")
            ws = wb.active
            encontrado = False
            for row in ws.iter_rows(min_row=2, values_only=False):
                if row[1].value and row[1].value.strip().lower() == nome:
                    telefone_entry.delete(0, tk.END)
                    telefone_entry.insert(0, row[4].value or "")
                    email_entry.delete(0, tk.END)
                    email_entry.insert(0, row[5].value or "")
                    encontrado = True
                    break
            if not encontrado:
                messagebox.showwarning("Aviso", "Contato não encontrado!")
        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao buscar contato.\n{e}")

    def salvar_alteracao():
        nome = nome_entry.get().strip().lower()
        novo_telefone = telefone_entry.get().strip()
        novo_email = email_entry.get().strip()
        try:
            wb = load_workbook("dados_brutos.xlsx")
            ws = wb.active
            alterado = False
            for row in ws.iter_rows(min_row=2, values_only=False):
                if row[1].value and row[1].value.strip().lower() == nome:
                    row[4].value = novo_telefone
                    row[5].value = novo_email
                    wb.save("dados_brutos.xlsx")
                    messagebox.showinfo("Sucesso", "Contato alterado com sucesso!")
                    form.destroy()
                    alterado = True
                    break
            if not alterado:
                messagebox.showwarning("Aviso", "Contato não encontrado!")
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível alterar o contato.\n{e}")

    tk.Button(form, text="Buscar", bg="#3399ff", fg="white", command=buscar_contato).grid(row=3, column=0, pady=20)
    tk.Button(form, text="Salvar Alterações", bg="#00cc66", fg="white", command=salvar_alteracao).grid(row=3, column=1, pady=20)
