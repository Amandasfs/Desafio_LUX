import tkinter as tk
from tkinter import messagebox
from openpyxl import load_workbook

# COMPONENTS
from ui.components.header import create_header as Header
from ui.components.sidebar import create_sidebar as Sidebar
from ui.components.toolbar import create_toolbar as Toolbar
from ui.components.cards import create_cards as Cards

# SCREENS
from .contatos_screen import ContatosScreen
from .empresas_screen import EmpresasScreen
from .consumo_screen import ConsumoScreen
from .graficos_screen import GraficosScreen
from .gestores_screen import GestoresScreen


class DashboardWindow:
    def __init__(self, root):
        self.root = root
        self.root.title("Dashboard - Gerenciamento")
        self.root.geometry("1000x700")

        # === Header ===
        self.header = Header(self.root)
        self.header.pack(side="top", fill="x")

        # === Funções de busca e organização atribuídas antes da toolbar ===
        self.buscar = self.buscar_dados
        self.organizar_dados = self.organizar_dados_excel

        # === Toolbar ===
        self.toolbar = Toolbar(self.root, self)
        self.toolbar.pack(side="top", fill="x")

        # === Frame principal para sidebar e conteúdo ===
        self.main_container = tk.Frame(self.root, bg="#f0f0f0")
        self.main_container.pack(side="top", expand=True, fill="both")

        # === Cria métodos de navegação ===
        self.create_navigation_methods()

        # === Sidebar ===
        self.sidebar = Sidebar(self.main_container, self)
        self.sidebar.pack(side="left", fill="y", padx=5, pady=5)

        # === Frame direito (Cards + Conteúdo dinâmico) ===
        self.right_content = tk.Frame(self.main_container, bg="#f0f0f0")
        self.right_content.pack(side="right", expand=True, fill="both", padx=5, pady=5)

        # === Cards ===
        self.cards_frame = tk.Frame(self.right_content, bg="#f0f0f0")
        self.cards_frame.pack(side="top", fill="x", padx=15, pady=(10, 5))
        self.cards = Cards(self.cards_frame)

        # === Conteúdo dinâmico abaixo dos cards ===
        self.main_content_frame = tk.Frame(self.right_content, bg="#ffffff", relief="solid", bd=1)
        self.main_content_frame.pack(side="top", fill="both", expand=True, padx=15, pady=5)

        # Inicializa a tela inicial
        self.show_screen("Home")

        # Debug layout
        self.root.after(100, self.debug_layout)

    # === Funções principais ===
    def buscar_dados(self):
        termo = getattr(self, "get_search_term", lambda: "")().lower().strip()
        print(f"Buscando por: {termo}")
        try:
            wb = load_workbook("dados_brutos.xlsx")
            ws = wb.active
            resultados = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if termo in str(row[1]).lower() or termo in str(row[3]).lower():
                    resultados.append(row)
            self.show_resultados(resultados)
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível buscar dados.\n{e}")

    def organizar_dados_excel(self):
        def executar_organizacao(tipo):
            try:
                wb = load_workbook("dados_brutos.xlsx")
                ws = wb.active
                dados = list(ws.iter_rows(min_row=2, values_only=True))

                # Escolhe a coluna de ordenação conforme o tipo
                if tipo == "Empresa":
                    dados.sort(key=lambda x: str(x[3]).strip().lower() if x[3] else "")
                elif tipo == "Colaborador":
                    dados.sort(key=lambda x: str(x[1]).strip().lower() if x[1] else "")
                elif tipo == "Ordem alfabética":
                    dados.sort(key=lambda x: str(x[1]).strip().lower() if x[1] else "")
                else:
                    messagebox.showwarning("Aviso", "Seleção inválida.")
                    return

                # Reescreve os dados ordenados
                for i, row in enumerate(dados, start=2):
                    for j, valor in enumerate(row, start=1):
                        ws.cell(row=i, column=j, value=valor)

                wb.save("dados_brutos.xlsx")

                messagebox.showinfo("Sucesso", f"Dados organizados por {tipo} com sucesso!")
                self.show_resultados(dados)

            except Exception as e:
                messagebox.showerror("Erro", f"Não foi possível organizar os dados.\n{e}")
            finally:
                popup.destroy()

        # === Popup de escolha ===
        popup = tk.Toplevel(self.root)
        popup.title("Organizar Dados")
        popup.geometry("300x220")
        popup.configure(bg="#f5f5f5")
        popup.resizable(False, False)

        tk.Label(
            popup,
            text="Organizar dados por:",
            font=("Arial", 12, "bold"),
            bg="#f5f5f5",
            fg="#00215b"
        ).pack(pady=15)

        botoes = [
            ("Empresa", "#004aad"),
            ("Colaborador", "#007bff"),
            ("Ordem alfabética", "#0099ff"),
        ]

        for nome, cor in botoes:
            tk.Button(
                popup,
                text=nome,
                width=20,
                bg=cor,
                fg="white",
                font=("Arial", 10, "bold"),
                command=lambda n=nome: executar_organizacao(n)
            ).pack(pady=5)

    def show_resultados(self, resultados):
        # Limpa conteúdo atual
        for widget in self.main_content_frame.winfo_children():
            widget.destroy()

        if not resultados:
            tk.Label(self.main_content_frame, text="Nenhum resultado encontrado.", bg="#ffffff").pack(pady=50)
            return

        # Mostra resultados em tabela simples
        for i, row in enumerate(resultados):
            for j, valor in enumerate(row):
                tk.Label(self.main_content_frame, text=valor, borderwidth=1, relief="solid", width=15).grid(row=i, column=j)

    # === Layout e navegação ===
    def debug_layout(self):
        print("=== DEBUG LAYOUT ===")
        print(f"Header: {self.header.winfo_ismapped()}")
        print(f"Toolbar: {self.toolbar.winfo_ismapped()}")
        print(f"Sidebar: {self.sidebar.winfo_ismapped()}")
        print(f"Cards Frame: {self.cards_frame.winfo_ismapped()}")
        print(f"Main Content: {self.main_content_frame.winfo_ismapped()}")
        print("====================")

    def create_navigation_methods(self):
        self.mostrar_contatos = lambda: self.show_screen("Contatos")
        self.mostrar_empresas = lambda: self.show_screen("Empresas")
        self.mostrar_consumo = lambda: self.show_screen("Consumo")
        self.mostrar_graficos = lambda: self.show_screen("Graficos")
        self.mostrar_gestores = lambda: self.show_screen("Gestores")

    def show_screen(self, screen_name):
        for widget in self.main_content_frame.winfo_children():
            widget.destroy()

        if screen_name == "Home":
            tk.Label(self.main_content_frame, text="Bem-vindo ao Dashboard!", font=("Arial", 20), bg="#ffffff").pack(pady=50)
        elif screen_name == "Contatos":
            ContatosScreen(self.main_content_frame)
        elif screen_name == "Empresas":
            EmpresasScreen(self.main_content_frame)
        elif screen_name == "Consumo":
            ConsumoScreen(self.main_content_frame)
        elif screen_name == "Graficos":
            GraficosScreen(self.main_content_frame)
        elif screen_name == "Gestores":
            GestoresScreen(self.main_content_frame)
        else:
            tk.Label(self.main_content_frame, text=f"Tela: {screen_name}", font=("Arial", 18), bg="#ffffff").pack(pady=50)
